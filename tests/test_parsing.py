import io
import unittest
from pathlib import Path

from tests import conftest  # noqa: F401
from skillmap.parsing import MAX_FILE_BYTES, ParseError, clean_text, extract_text

EXAMPLES = Path(__file__).resolve().parent.parent / "data" / "sample_resumes" / "upload_examples"


class ParsingTests(unittest.TestCase):
    def test_txt_utf8_and_cp1252(self):
        self.assertIn("Résumé", extract_text("cv.txt", "Résumé\nPython".encode("utf-8")))
        self.assertIn("café", extract_text("cv.txt", "café".encode("cp1252")))

    def test_docx_round_trip(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("Backend engineer")
        doc.add_paragraph("Skills: Python, Docker")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Cloud"
        table.rows[0].cells[1].text = "AWS"
        buffer = io.BytesIO()
        doc.save(buffer)
        text = extract_text("cv.docx", buffer.getvalue())
        for expected in ("Backend engineer", "Python, Docker", "Cloud | AWS"):
            self.assertIn(expected, text)

    def test_sample_docx_file(self):
        path = EXAMPLES / "data_analyst.docx"
        text = extract_text(path.name, path.read_bytes())
        self.assertIn("Tableau", text)

    def test_sample_pdf_file(self):
        path = EXAMPLES / "backend_engineer.pdf"
        text = extract_text(path.name, path.read_bytes())
        self.assertIn("PostgreSQL", text)
        self.assertIn("Django", text)

    def test_pdf_without_text_is_reported(self):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)
        with self.assertRaises(ParseError) as ctx:
            extract_text("scan.pdf", buffer.getvalue())
        self.assertIn("scanned", str(ctx.exception))

    def test_corrupt_files(self):
        with self.assertRaises(ParseError):
            extract_text("bad.pdf", b"%PDF-1.4 definitely not a pdf")
        with self.assertRaises(ParseError):
            extract_text("bad.docx", b"not a zip file")

    def test_unsupported_empty_and_oversized(self):
        with self.assertRaises(ParseError):
            extract_text("cv.doc", b"data")
        with self.assertRaises(ParseError):
            extract_text("cv.txt", b"")
        with self.assertRaises(ParseError):
            extract_text("cv.txt", b"a" * (MAX_FILE_BYTES + 1))

    def test_clean_text(self):
        self.assertEqual(clean_text("a  b\r\n\r\n\r\n\r\nc \n"), "a b\n\nc")


if __name__ == "__main__":
    unittest.main()
