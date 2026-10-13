"""Extract plain text from uploaded resume files (PDF, DOCX, TXT)."""

from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")
MAX_FILE_BYTES = 5 * 1024 * 1024  # 5 MB


class ParseError(Exception):
    """Raised with a user-readable message when a file cannot be read."""


def clean_text(text: str) -> str:
    """Normalise line endings and whitespace without changing the content."""
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ").replace("\u200b", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_txt(data: bytes) -> str:
    # UTF-16 only when a BOM says so: it "decodes" almost any even-length bytes.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    raise ParseError("The text file could not be decoded. Save it as UTF-8 and try again.")


def _read_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
        from pypdf.errors import PyPdfError
    except ImportError as exc:  # pragma: no cover - dependency missing
        raise ParseError("PDF support needs the 'pypdf' package (pip install pypdf).") from exc

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                if not reader.decrypt(""):
                    raise ParseError("This PDF is password-protected. Remove the password and upload it again.")
            except ParseError:
                raise
            except Exception as exc:  # noqa: BLE001 - pypdf raises various errors here
                raise ParseError("This PDF is encrypted and could not be opened.") from exc
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ParseError:
        raise
    except (PyPdfError, ValueError, OSError, KeyError) as exc:
        raise ParseError("This PDF could not be read. It may be damaged; try exporting it again.") from exc

    text = "\n".join(pages)
    if not text.strip():
        raise ParseError(
            "No selectable text was found in this PDF. It may be a scanned image. "
            "OCR is not supported, so export a text-based PDF or paste the text instead."
        )
    return text


def _read_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dependency missing
        raise ParseError("DOCX support needs the 'python-docx' package (pip install python-docx).") from exc

    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - python-docx raises several error types
        raise ParseError("This DOCX file could not be read. Make sure it is a real .docx, not .doc.") from exc

    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            seen: set[int] = set()
            cells = []
            for cell in row.cells:
                if id(cell._tc) in seen:  # merged cells repeat
                    continue
                seen.add(id(cell._tc))
                cells.append(cell.text.strip())
            lines.append(" | ".join(c for c in cells if c))
    text = "\n".join(lines)
    if not text.strip():
        raise ParseError("This DOCX file has no readable text.")
    return text


def extract_text(filename: str, data: bytes) -> str:
    """Return cleaned text from a resume file.

    Raises ``ParseError`` with a message that can be shown to the user.
    """
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        shown = suffix or "no extension"
        raise ParseError(f"Unsupported file type ({shown}). Upload a PDF, DOCX, or TXT file.")
    if not data:
        raise ParseError("The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise ParseError(f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")

    if suffix == ".pdf":
        raw = _read_pdf(data)
    elif suffix == ".docx":
        raw = _read_docx(data)
    else:
        raw = _read_txt(data)

    text = clean_text(raw)
    if not text:
        raise ParseError("No text could be extracted from this file.")
    return text
