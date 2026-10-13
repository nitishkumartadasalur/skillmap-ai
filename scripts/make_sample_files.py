"""Regenerate the PDF and DOCX upload examples from the sample .txt resumes.

Needs reportlab (see requirements-dev.txt) for the PDF. Run from the project root:
    python scripts/make_sample_files.py
"""

from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "sample_resumes"
OUT = SRC / "upload_examples"
OUT.mkdir(exist_ok=True)


def make_docx(src: Path, dest: Path) -> None:
    doc = Document()
    for line in src.read_text(encoding="utf-8").splitlines():
        if line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line.isupper() and line.strip():
            doc.add_heading(line.title(), level=2)
        else:
            doc.add_paragraph(line)
    doc.save(dest)


def make_pdf(src: Path, dest: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = []
    for line in src.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            story.append(Spacer(1, 6))
        elif line.startswith("- "):
            story.append(Paragraph("\u2022 " + escape(line[2:]), styles["BodyText"]))
        elif line.isupper():
            story.append(Paragraph(f"<b>{escape(line)}</b>", styles["Heading3"]))
        else:
            story.append(Paragraph(escape(line), styles["BodyText"]))
    SimpleDocTemplate(str(dest), pagesize=A4).build(story)


if __name__ == "__main__":
    make_docx(SRC / "data_analyst.txt", OUT / "data_analyst.docx")
    make_pdf(SRC / "backend_engineer.txt", OUT / "backend_engineer.pdf")
    print("Wrote", *sorted(p.name for p in OUT.iterdir()))
